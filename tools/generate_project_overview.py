from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "BrewMart_Project_Overview.docx"

COFFEE = "6B3F2A"
INK = "24313A"
MUTED = "66737C"
CREAM = "F7F2EB"
LINE = "D9CEC2"
ACCENT = "B46B3C"


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def set_cell_text(cell, text, bold=False, color=INK, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=9)
        shade(table.rows[0].cells[index], COFFEE)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value), size=8.5)
            shade(cells[index], "FFFFFF" if len(table.rows) % 2 else CREAM)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(document, items, level=0):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.22 + (level * 0.18))
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(item)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_callout(document, title, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F1E5D8")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COFFEE)
    paragraph.add_run("\n" + text)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size, color in (("Title", 28, COFFEE), ("Heading 1", 16, COFFEE), ("Heading 2", 11, ACCENT)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("BrewMart Project Overview  |  Prepared for project review")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)


def build():
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("BrewMart")
    run.font.name = "Aptos Display"
    run.font.size = Pt(30)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COFFEE)
    subtitle = document.add_paragraph("Project overview and presentation brief")
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)
    metadata = document.add_paragraph(f"Django coffee e-commerce platform  |  Review edition: {date.today().isoformat()}")
    metadata.paragraph_format.space_after = Pt(12)
    metadata.runs[0].font.size = Pt(8.5)
    metadata.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(document, "One-sentence description", "BrewMart is a multi-role coffee e-commerce web application that connects customer shopping, seller order preparation, delivery fulfillment, and administrator operations in one Django system.")

    add_heading(document, "1. Executive Summary")
    document.add_paragraph("This project is more than a product catalogue. It is a complete commerce workflow: customers discover products, manage carts and wishlists, check out, receive order updates, and review purchases; sellers manage assigned sales; delivery users fulfil orders; and administrators manage the business through dashboards and Django admin.")
    add_bullets(document, [
        "Primary business domain: online coffee retail and order fulfilment.",
        "Application style: server-rendered Django web application with a database-backed domain model.",
        "Currency and regional context: Nepalese rupees, with Asia/Kathmandu timezone support.",
        "Current maturity: feature-rich local/development system with deployment foundations, but payment and wallet functions remain demonstrations.",
    ])

    add_heading(document, "2. What Type of System Is It?")
    add_table(document, ["Dimension", "Implemented meaning"], [
        ("System category", "Role-based B2C e-commerce and fulfilment management system"),
        ("Users", "Buyers, sellers, delivery personnel, administrators, and Django superusers"),
        ("Core transaction", "Product selection -> cart -> checkout/payment -> order -> fulfilment -> buyer confirmation"),
        ("Architecture", "Modular Django monolith with separated domain apps and shared templates"),
        ("Data style", "Relational models with migrations, order snapshots, and transactional updates"),
    ], widths=[1.45, 5.9])

    add_heading(document, "3. Technology Stack")
    add_table(document, ["Layer", "Technology / implementation"], [
        ("Language", "Python 3.10+"),
        ("Backend", "Django 5.1, Django authentication, forms, ORM, sessions, admin, middleware"),
        ("Frontend", "Django templates, HTML, CSS, JavaScript, responsive shared layout"),
        ("Database", "SQLite by default for development; PostgreSQL through DATABASE_URL for deployment"),
        ("Images", "Pillow and Django ImageField uploads under media/"),
        ("Static files", "WhiteNoise compressed manifest storage"),
        ("Configuration", "python-dotenv and environment variables from .env"),
        ("Server/deployment", "Django WSGI entry point and Gunicorn dependency; reverse-proxy HTTPS guidance"),
        ("Reporting/support", "ReportLab dependency for PDF-related functionality and management commands"),
    ], widths=[1.45, 5.9])

    add_heading(document, "4. Project Structure")
    add_table(document, ["Module", "Responsibility"], [
        ("config", "Settings, root URL routing, WSGI, database and security configuration"),
        ("core", "Home page, shared context, SEO sitemap and robots response"),
        ("accounts", "Registration, profiles, roles, wallets, and wallet transfers"),
        ("products", "Categories, product catalogue, pricing, stock, SEO data, recommendations"),
        ("cart", "Authenticated and anonymous carts, cart items, quantities, totals"),
        ("wishlist", "User saved products"),
        ("orders", "Checkout, payment simulators, order creation, history, and workflow"),
        ("notifications", "In-app notifications and optional email/WhatsApp delivery"),
        ("reviews", "Product ratings and customer comments"),
        ("dashboard", "Admin, seller, delivery operations, analytics, and notification testing"),
        ("templates / static / media", "Presentation templates, CSS/JavaScript, and uploaded images"),
    ], widths=[1.65, 5.7])

    add_heading(document, "5. User Roles and Capabilities")
    add_table(document, ["Role", "Main capabilities"], [
        ("Buyer", "Register, manage profile, browse products, cart, wishlist, reviews, checkout, wallet transfers, order history, and delivery confirmation"),
        ("Seller", "Manage owned products, view assigned orders, accept and prepare orders, make orders ready, assign a rider, or cancel eligible orders"),
        ("Delivery person", "View available deliveries, claim work, pick up orders, mark them out for delivery, and mark them delivered"),
        ("Admin", "Manage users, roles, products, categories, orders, notification settings, analytics, and operational records"),
    ], widths=[1.45, 5.9])

    add_heading(document, "6. Main Customer and Business Workflows")
    add_heading(document, "Customer purchase flow", level=2)
    add_bullets(document, [
        "Browse home, catalogue, categories, product details, related products, and recommendations.",
        "Add products to a persistent user cart or a session cart, then adjust quantities or remove items.",
        "Proceed through authenticated checkout, delivery information, and a selected payment method.",
        "On successful demo payment, the system creates the order, snapshots item data, reduces stock, clears the cart, and records transaction proof.",
        "Track the order and confirm receipt after delivery; submit one review per product.",
    ])
    add_heading(document, "Order state machine", level=2)
    document.add_paragraph("PLACED -> ACCEPTED -> PREPARING -> READY_FOR_DELIVERY -> ASSIGNED -> PICKED_UP -> OUT_FOR_DELIVERY -> DELIVERED -> COMPLETED")
    document.add_paragraph("Cancellation is available to the assigned seller from PLACED, ACCEPTED, or PREPARING. The seller may also directly assign a delivery rider while preparing; otherwise a delivery user claims a ready order.")
    add_heading(document, "Supporting commerce features", level=2)
    add_bullets(document, [
        "Stored-value wallet transfers run atomically, lock both wallets, reject self-transfers and insufficient balances, and create a unique transfer reference.",
        "Notifications are created in-app for order events, with optional email and WhatsApp channels configured through environment variables.",
        "Signed-in users receive collaborative recommendations based on co-purchases; visitors receive trending or newest-product fallbacks.",
    ])

    add_heading(document, "7. Core Data Model")
    add_table(document, ["Entity", "Important data represented"], [
        ("Profile", "User role, phone, address, city, and postal code"),
        ("Wallet / WalletTransfer", "Non-negative balance and auditable peer-to-peer transfer records"),
        ("Category / Product", "Catalog hierarchy, price/discount, stock, SKU, seller, brand, image, active/featured status"),
        ("Cart / CartItem", "User or session ownership, product quantity, unique product per cart"),
        ("Wishlist / Review", "Saved products and one user review per product with rating/comment"),
        ("Order / OrderItem", "Buyer, seller, delivery person, delivery snapshot, payment proof, status timestamps, item price/name snapshots"),
        ("Notification", "Recipient, order reference, message, type, read state, and timestamp"),
    ], widths=[1.65, 5.7])
    document.add_paragraph("Order and item snapshots preserve historical accuracy even if a product name, price, or customer delivery profile changes later.")

    add_heading(document, "8. Security and Reliability")
    add_bullets(document, [
        "Django authentication, password validators, login protection, CSRF middleware, role decorators, and ownership checks protect user actions.",
        "Order creation, stock changes, workflow changes, and wallet transfers use database transactions where multiple records must remain consistent.",
        "Transaction payloads are canonicalized and protected with SHA-256 hashing plus HMAC-SHA256 signing using SECRET_KEY; a reusable verifier is implemented.",
        "When DEBUG=False, secure cookies, HTTPS redirect support, proxy-aware HTTPS, content-type sniffing protection, referrer policy, and configurable HSTS are enabled.",
        "Secrets and provider credentials are expected in environment variables and must not be copied into reports, source control, or logs.",
    ])
    add_callout(document, "Important boundary", "Card, Khalti, and eSewa classes are local simulators. They do not contact banks or payment providers. Before production, use an official provider flow with a pending order and server-side callback or lookup verification. The wallet is also demonstration stored value, not a real financial service.")

    add_heading(document, "9. Marketing, Analytics, and Discoverability")
    add_bullets(document, [
        "Optional Google Analytics 4 and Meta/Facebook Pixel support controlled by environment variables.",
        "Product impression, product selection, product view, add-to-cart, and purchase events are wired through the shared frontend analytics script.",
        "Optional Google AdSense script and reusable ad slot; the current ad slot ID is a placeholder until configured.",
        "SEO support includes page metadata, canonical URLs, Open Graph/Twitter tags, product schema.org JSON-LD, sitemap.xml, robots.txt, and image alt text.",
        "Recommendation logic uses collaborative filtering first, then same-category products, then trending products as fallbacks.",
    ])

    add_heading(document, "10. Running and Demonstrating the Project")
    document.add_paragraph("Typical local sequence from the project root:")
    add_bullets(document, [
        "Create/activate the Python virtual environment and install requirements.txt.",
        "Run python manage.py migrate.",
        "Create an administrator with python manage.py createsuperuser when needed.",
        "Run python manage.py check, then python manage.py runserver.",
        "Open the store at http://127.0.0.1:8000/, admin at /admin/, and the operational dashboard at /dashboard/.",
    ])
    add_heading(document, "Suggested presentation demonstration", level=2)
    add_bullets(document, [
        "Show the public catalogue and product detail page.",
        "Create a buyer account, add a product to the cart, and demonstrate checkout with clearly labelled demo payment behavior.",
        "Show the resulting order and notification, then use seller and delivery roles to advance it through the state machine.",
        "Open the admin dashboard to show product, user, order, low-stock, and revenue management.",
        "Close with the architecture, security boundary, and production work still required.",
    ])

    add_heading(document, "11. Current Status and Next Steps")
    add_table(document, ["Area", "Current status / interpretation"], [
        ("Implemented", "Core shopping, accounts, catalogue, cart, wishlist, reviews, checkout, order fulfilment, dashboards, notifications, wallet transfers, SEO, analytics hooks, and recommendations"),
        ("Development ready", "Django system check passes; SQLite and local HTTP/HTTPS workflows are documented"),
        ("Production preparation", "Configure PostgreSQL, real secrets, domain/CSRF settings, trusted HTTPS, media storage, email/WhatsApp providers, and deployment server"),
        ("Required before live payments", "Replace simulators with official gateways and verify provider callbacks server-side before finalizing orders"),
        ("Quality gap", "No automated test modules are currently visible; add focused tests for checkout, stock, permissions, wallet concurrency, and workflow transitions"),
    ], widths=[1.65, 5.7])

    add_heading(document, "12. Presentation Summary")
    document.add_paragraph("BrewMart can be presented as a modular, role-based Django e-commerce and fulfilment management system for a coffee business. Its distinguishing strength is the complete operational loop: it does not stop at shopping and checkout, but models seller responsibility, delivery assignment, status transitions, notifications, administration, and customer confirmation. The project is suitable for academic demonstration or further product development, provided the clearly identified demo integrations are replaced before real-world financial use.")

    add_heading(document, "Reference Files")
    document.add_paragraph("The implementation details summarized here are based on the repository source and project documentation, especially AI_PROJECT_CONTEXT.md, PROJECT_MEMORY.md, docs/RUN_PROJECT.md, config/settings.py, config/urls.py, accounts/, products/, cart/, orders/, notifications/, reviews/, and dashboard/.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()